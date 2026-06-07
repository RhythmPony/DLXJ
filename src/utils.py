import configparser
import logging
import os
import re
import sqlite3

import cv2
import numpy as np
import xlsxwriter
from PyQt5.QtCore import Qt
from PyQt5.QtGui import QPixmap, QImage, qRgb
from docx import Document
from docx.shared import Cm

from src.wordTemplate import generate_word_with_template


def write_docx(doc_export_path: str, filename: str, records: list, queue):
    full_path = os.path.join(doc_export_path, filename)
    try:
        document = Document(full_path)
    except Exception as e:
        document = Document()

    for record in records:
        generate_word_with_template(document, record)
        queue.put(os.path.splitext(filename)[0])
    queue.put(os.path.splitext(filename)[0] + "_end")

    sections = document.sections

    width, height = get_page_size()
    margin_top, margin_right, margin_bottom, margin_left = get_page_margin()

    for section in sections:
        section.page_width = Cm(width)
        section.page_height = Cm(height)
        section.top_margin = Cm(margin_top)
        section.bottom_margin = Cm(margin_bottom)
        section.left_margin = Cm(margin_left)
        section.right_margin = Cm(margin_right)

    if not os.path.exists(doc_export_path):
        os.makedirs(doc_export_path)
    document.save(full_path)


def write_xlsx(db_name, xlsx_name):
    table_names = get_table_names(db_name)

    workbook = xlsxwriter.Workbook(xlsx_name)
    for table_name in table_names:
        if table_name != 'sqlite_sequence':
            records = get_table_records(db_name, table_name)
            worksheet = workbook.add_worksheet(table_name)
            title = get_table_headers(db_name, table_name)
            title_format = workbook.add_format({'bold': True, 'align': 'center', 'font_size': 13})
            body_format = workbook.add_format({'align': 'center', 'font_size': 11})
            worksheet.write_row(0, 0, title, title_format)
            row = 1
            for record in records:
                worksheet.write_row(row, 0, record, body_format)
                row += 1
    workbook.close()
    return 200


def get_page_size():
    try:
        size_hint_config = get_option("sys", "word_page_size", "21,29.7")
        size_hint = size_hint_config.split(",")
        size_hint = [float(i) for i in size_hint]
        try:
            width = size_hint[0]
            height = size_hint[1]
        except:
            width = 21
            height = 29.7
        return width, height
    except Exception as e:
        logging.warning(str(e))
        return 21, 29.7


def get_page_margin():
    try:
        margin_config = get_option("sys", "word_page_margin", "2.54,1.91,2.54,1.91")
        margin = margin_config.split(",")
        margin = [float(i) for i in margin]
        if len(margin) == 2:
            margin_top = margin[0]
            margin_right = margin[1]
            margin_bottom = margin[0]
            margin_left = margin[1]
        elif len(margin) == 4:
            margin_top = margin[0]
            margin_right = margin[1]
            margin_bottom = margin[2]
            margin_left = margin[3]
        else:
            margin_top = 2.54
            margin_right = 1.91
            margin_bottom = 2.54
            margin_left = 1.91
        return margin_top, margin_right, margin_bottom, margin_left
    except Exception as e:
        logging.warning(str(e))
        return 2.54, 1.91, 2.54, 1.91


def read_sys_config(sys_config_path: str):
    try:
        config = configparser.ConfigParser()
        config.read(sys_config_path, encoding='utf-8')
        return config
    except Exception as e:
        logging.error(str(e))
        return -1


def read_options(defect_option_path: str):
    """
    读取配置选项并以字典形式返回
    :return: options
    """""
    try:
        temp_dic = {}
        temp_str = ''
        with open(defect_option_path, 'r+', encoding='utf-8') as settings:
            for line in settings.readlines():
                if line.strip():
                    if not line.strip().startswith('#'):
                        # 读取方括号中的选项标题保存为字典的键
                        if re.search(r'\[\w*]', line.strip()):
                            temp_dic.update({re.search(r'\[\w+]', line.strip())[0]: []})
                            temp_str = re.search(r'\[\w+]', line.strip())[0]

                        # 读取配置选项列表保存为字典的值
                        else:
                            temp_list = temp_dic[temp_str]
                            temp_list.append(line.strip())
                            temp_dic.update({temp_str: temp_list})
        return temp_dic

    except Exception as e:
        logging.error(str(e))
        return -1


def get_grid_position(num: int):
    if num % 2 == 0:
        row = num / 2
        column = 0
    else:
        row = (num - 1) / 2
        column = 1
    return row, column


def remove_black_border(image: np.ndarray):
    img = cv2.medianBlur(image, 5)  # 中值滤波，去除黑色边际中可能含有的噪声干扰
    b = cv2.threshold(img, 3, 255, cv2.THRESH_BINARY)  # 调整裁剪效果
    binary_image = b[1]  # 二值图--具有三通道
    binary_image = cv2.cvtColor(binary_image, cv2.COLOR_BGR2GRAY)
    # print(binary_image.shape)     #改为单通道

    edges_y, edges_x = np.where(binary_image == 255)  # h, w
    bottom = min(edges_y)
    top = max(edges_y)

    left = min(edges_x)
    right = max(edges_x)
    height = top - bottom
    width = right - left

    res_image = image[bottom:bottom + height, left:left + width]

    return res_image


def q_image_2_cv2(q_image: QImage):
    width = q_image.width()
    height = q_image.height()

    ptr = q_image.bits()
    ptr.setsize(q_image.byteCount())
    cv_image = np.array(ptr).reshape(height, width, 4)

    return cv_image


def cv2_2_q_image(cv_image):
    width = cv_image.shape[1]  # 获取图片宽度
    height = cv_image.shape[0]  # 获取图片高度

    pixmap = QPixmap(width, height)  # 根据已知的高度和宽度新建一个空的QPixmap,
    q_image = pixmap.toImage()  # 将pximap转换为QImage类型的q_image

    # 循环读取cv_image的每个像素的r,g,b值，构成qRgb对象，再设置为q_image内指定位置的像素
    for row in range(height):
        for col in range(width):
            b = cv_image[row, col, 0]
            g = cv_image[row, col, 1]
            r = cv_image[row, col, 2]

            pix = qRgb(r, g, b)
            q_image.setPixel(col, row, pix)

    return q_image  # 转换完成，返回


def resize_pixmap(pixmap: QPixmap, ratio: list = None):
    image_resize_hint_list = [None, Qt.IgnoreAspectRatio, Qt.KeepAspectRatio, Qt.KeepAspectRatioByExpanding]
    hint = int(get_option("sys", "image_resize_hint", 1))
    image_resize_hint = image_resize_hint_list[hint]

    if not image_resize_hint:
        return pixmap
    row_width = pixmap.size().width()
    row_height = pixmap.size().height()

    if row_height / row_width <= ratio[1] / ratio[0]:
        target_height = int(row_height)
        target_width = int(target_height * ratio[0] / ratio[1])
        target = pixmap.scaled(target_width, target_height, image_resize_hint)
        return target

    else:
        target_width = int(row_width)
        target_height = int(target_width * ratio[1] / ratio[0])
        target = pixmap.scaled(target_width, target_height, image_resize_hint)
        return target


def get_option(section, key, default):
    from base import BASE_DIR, sys_config_path
    config = read_sys_config(os.path.join(BASE_DIR, sys_config_path))
    try:
        value = config.get(section, key)
        if value == "default":
            value = default
    except Exception as e:
        logging.warning(str(e) + "--utils.get_option")
        value = default
    return value


def reset_option(section, key, value):
    from base import BASE_DIR, sys_config_path
    config = read_sys_config(os.path.join(BASE_DIR, sys_config_path))
    config.set(section, key, value)
    with open(os.path.join(BASE_DIR, sys_config_path), 'w', encoding='utf-8') as configfile:
        config.write(configfile)


def get_table_names(db_name):
    if os.path.isfile(os.path.abspath(db_name)):
        try:
            conn = sqlite3.connect(db_name)
            cursor = conn.cursor()
            table_names = [i[0] for i in
                           cursor.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")]
            conn.close()
            return table_names
        except Exception as e:
            logging.warning(str(e))


def get_table_creations(db_name):
    if os.path.isfile(os.path.abspath(db_name)):
        try:
            conn = sqlite3.connect(db_name)
            cursor = conn.cursor()
            cursor.execute("SELECT name,sql FROM sqlite_master WHERE type='table'")
            table_info_list = [(i[0], i[1]) for i in cursor.fetchall() if "sqlite_sequence" not in i[0]]
            conn.close()
            return table_info_list
        except Exception as e:
            logging.warning(str(e))


def get_table_headers(db_name, table_name):
    conn = sqlite3.connect(db_name)
    cursor = conn.cursor()
    sql = f"PRAGMA  table_info({table_name})"
    cursor.execute(sql)
    headers = [i[1] for i in cursor.fetchall()]
    conn.close()
    return headers


def get_table_records(db_name, table_name):
    conn = sqlite3.connect(db_name)
    cursor = conn.cursor()
    sql = f"SELECT * FROM {table_name}"
    cursor.execute(sql)
    records = cursor.fetchall()
    conn.close()
    return records


if __name__ == '__main__':
    pass
