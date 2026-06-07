import os

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def generate_word_with_template(document: Document, record: tuple):
    marker = "level"
    try:
        if record[9] == 1:
            title = marker + str(record[0]) + '.' + record[1] + record[2]
            Head = document.add_heading('', level=2)
            run = Head.add_run(title)
            run.font.name = '仿宋'
            run.font.size = Pt(14)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')
            run.font.color.rgb = RGBColor(0, 0, 0)

            single_column_table = document.add_table(5, 4, document.styles['Table Grid'])
            single_column_table.cell(0, 0).merge(single_column_table.cell(0, 1))

            run = single_column_table.cell(0, 0).paragraphs[0].add_run(u'线路名称:' + record[1])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            run = single_column_table.cell(0, 2).paragraphs[0].add_run(u'杆塔编号:' + record[2])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            run = single_column_table.cell(0, 3).paragraphs[0].add_run(u'巡检方式：' + record[3])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            single_column_table.cell(1, 0).merge(single_column_table.cell(1, 1))
            run = single_column_table.cell(1, 0).paragraphs[0].add_run(u'缺陷部件：' + record[4])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            run = single_column_table.cell(1, 2).paragraphs[0].add_run(u'缺陷等级：' + record[5])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            run = single_column_table.cell(1, 3).paragraphs[0].add_run(u'巡检日期：' + record[6])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            single_column_table.cell(2, 0).merge(single_column_table.cell(2, 1))
            run = single_column_table.cell(2, 0).paragraphs[0].add_run(u'缺陷类型：' + record[7])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            single_column_table.cell(2, 2).merge(single_column_table.cell(2, 3))
            run = single_column_table.cell(2, 2).paragraphs[0].add_run(u'缺陷描述：' + record[8])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            single_column_table.cell(3, 0).merge(single_column_table.cell(3, 1)).merge(
                single_column_table.cell(3, 2).merge(single_column_table.cell(3, 3)))
            run = single_column_table.cell(3, 0).paragraphs[0].add_run(u'现场照片：' + record[10])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            single_column_table.cell(4, 0).merge(single_column_table.cell(4, 1)).merge(
                single_column_table.cell(4, 2).merge(single_column_table.cell(4, 3)))
            single_column_table.cell(4, 0).width = Cm(19.2)
            paragraph2 = single_column_table.cell(4, 0).paragraphs[0]
            paragraph2.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

            image_path = os.path.join(os.getcwd(), record[11])
            paragraph2.add_run().add_picture(image_path, width=Cm(13))

            for row in single_column_table.rows:
                row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

            document.add_paragraph('')
        if record[9] == 2:
            title = marker + str(record[0]) + '.' + record[1] + record[2]
            Head = document.add_heading('', level=2)
            run = Head.add_run(title)
            run.font.name = '仿宋'
            run.font.size = Pt(14)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')
            run.font.color.rgb = RGBColor(0, 0, 0)

            double_column_table = document.add_table(5, 4, document.styles['Table Grid'])
            double_column_table.cell(0, 0).merge(double_column_table.cell(0, 1))

            run = double_column_table.cell(0, 0).paragraphs[0].add_run(u'线路名称:' + record[1])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            double_column_table.cell(0, 2).paragraphs[0].add_run(u'杆塔编号:' + record[2])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            run = double_column_table.cell(0, 3).paragraphs[0].add_run(u'巡检方式：' + record[3])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            double_column_table.cell(1, 0).merge(double_column_table.cell(1, 1))
            run = double_column_table.cell(1, 0).paragraphs[0].add_run(u'缺陷部件：' + record[4])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            run = double_column_table.cell(1, 2).paragraphs[0].add_run(u'缺陷等级：' + record[5])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            run = double_column_table.cell(1, 3).paragraphs[0].add_run(u'巡检日期：' + record[6])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            double_column_table.cell(2, 0).merge(double_column_table.cell(2, 1))
            run = double_column_table.cell(2, 0).paragraphs[0].add_run(u'缺陷类型：' + record[7])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            double_column_table.cell(2, 2).merge(double_column_table.cell(2, 3))
            run = double_column_table.cell(2, 2).paragraphs[0].add_run(u'缺陷描述：' + record[8])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            double_column_table.cell(3, 0).merge(double_column_table.cell(3, 1)).merge(
                double_column_table.cell(3, 2).merge(double_column_table.cell(3, 3)))
            run = double_column_table.cell(3, 0).paragraphs[0].add_run(u'现场照片：' + record[10])
            run.font.name = '仿宋'
            run.font.size = Pt(12)
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

            double_column_table.cell(4, 0).merge(double_column_table.cell(4, 1))
            double_column_table.cell(4, 2).merge(double_column_table.cell(4, 3))
            double_column_table.cell(4, 0).width = Cm(9.1)
            double_column_table.cell(4, 2).width = Cm(9.1)
            double_column_table.cell(4, 0).height = Cm(9.0)
            paragraph1 = double_column_table.cell(4, 0).paragraphs[0]
            paragraph1.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

            image_path = os.path.join(os.getcwd(), record[11])
            paragraph1.add_run().add_picture(image_path, width=Cm(9))

            paragraph = double_column_table.cell(4, 2).paragraphs[0]
            paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

            image_path = os.path.join(os.getcwd(), record[12])
            paragraph.add_run().add_picture(image_path, width=Cm(9))

            for row in double_column_table.rows:
                row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

            document.add_paragraph('')
    except Exception as e:
        print(e)
        return -1
